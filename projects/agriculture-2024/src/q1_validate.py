"""Independent business-rule validation for solved Question 1 plans."""

from __future__ import annotations

import math
import re
import zipfile
from collections import defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from q1_data import (
    GREENHOUSE_TYPES,
    IRRIGATED_LAND,
    OPEN_LAND_TYPES,
    SEASON_FIRST,
    SEASON_SECOND,
    SEASON_SINGLE,
    YEARS,
    Q1Data,
    compatible_crop_ids,
    sha256,
)
from q1_model import PlantKey, Q1Solution, SalesKey


DEFAULT_TOLERANCE = 1e-6


class Q1ValidationError(RuntimeError):
    """Raised when a solved plan or filled template fails a hard check."""


@dataclass(frozen=True)
class ValidationCheck:
    name: str
    passed: bool
    max_violation: float
    details: str


@dataclass(frozen=True)
class ValidationReport:
    case_name: str
    passed: bool
    tolerance: float
    checks: tuple[ValidationCheck, ...]
    recomputed_profit_yuan: float
    solver_profit_yuan: float
    production_total_jin: float
    normal_sales_total_jin: float
    excess_total_jin: float

    def to_dict(self) -> dict[str, Any]:
        return {
            **{key: value for key, value in asdict(self).items() if key != "checks"},
            "checks": [asdict(check) for check in self.checks],
        }


def _check(name: str, violation: float, tolerance: float, details: str) -> ValidationCheck:
    return ValidationCheck(
        name=name,
        passed=violation <= tolerance,
        max_violation=max(0.0, float(violation)),
        details=details,
    )


def _production_by_sales_group(
    data: Q1Data, x: dict[PlantKey, float]
) -> dict[SalesKey, float]:
    production: dict[SalesKey, float] = defaultdict(float)
    for (plot, crop_id, year, season), area in x.items():
        land_type = data.plots[plot].land_type
        parameter = data.parameters[(crop_id, land_type, season)]
        production[(crop_id, land_type, year, season)] += (
            area * parameter.yield_jin_per_mu
        )
    return dict(production)


def validate_solution(
    solution: Q1Solution, *, tolerance: float = DEFAULT_TOLERANCE
) -> ValidationReport:
    data = solution.model.data
    x = solution.x
    z = solution.z
    regimes = solution.regimes
    normal = solution.normal_sales
    excess = solution.excess_sales
    checks: list[ValidationCheck] = []

    minimum_area = min(x.values()) if x else 0.0
    checks.append(
        _check(
            "nonnegative_planting_area",
            max(0.0, -minimum_area),
            tolerance,
            f"minimum raw area={minimum_area:.12g}",
        )
    )

    binary_values = list(z.values()) + list(regimes.values())
    integrality_violation = max(
        (abs(value - round(value)) for value in binary_values), default=0.0
    )
    checks.append(
        _check(
            "binary_integrality",
            integrality_violation,
            tolerance,
            f"checked {len(binary_values)} activation/regime binaries",
        )
    )

    linkage_violation = max(
        (
            area - data.plots[key[0]].area_mu * z[key]
            for key, area in x.items()
        ),
        default=0.0,
    )
    checks.append(
        _check(
            "area_activation_linkage",
            linkage_violation,
            tolerance,
            "x <= plot_area * z with plot-specific tight bounds",
        )
    )

    area_by_plot_season: dict[tuple[str, int, str], float] = defaultdict(float)
    for (plot, _crop, year, season), area in x.items():
        area_by_plot_season[(plot, year, season)] += area
    capacity_violation = max(
        (
            area_by_plot_season[(plot, year, season)] - data.plots[plot].area_mu
            for plot, year, season in data.available_plot_seasons
        ),
        default=0.0,
    )
    checks.append(
        _check(
            "plot_season_capacity",
            capacity_violation,
            tolerance,
            f"checked {len(data.available_plot_seasons)} plot-year-season capacities",
        )
    )

    incompatible_positive: list[PlantKey] = []
    for key, area in x.items():
        if area <= tolerance:
            continue
        plot, crop_id, _year, season = key
        allowed = compatible_crop_ids(data.plots[plot].land_type, season, data.crops)
        if crop_id not in allowed:
            incompatible_positive.append(key)
    checks.append(
        _check(
            "crop_land_season_compatibility",
            float(len(incompatible_positive)),
            0.0,
            f"incompatible positive cells={incompatible_positive[:10]}",
        )
    )

    greenhouse_violations: list[PlantKey] = []
    for key, area in x.items():
        if area <= tolerance:
            continue
        plot, crop_id, _year, season = key
        land_type = data.plots[plot].land_type
        crop = data.crops[crop_id]
        if land_type == "普通大棚":
            valid = (season == SEASON_FIRST and crop.is_vegetable and crop.name not in {"大白菜", "白萝卜", "红萝卜"}) or (
                season == SEASON_SECOND and crop.is_fungus
            )
        elif land_type == "智慧大棚":
            valid = season in {SEASON_FIRST, SEASON_SECOND} and crop.is_vegetable and crop.name not in {"大白菜", "白萝卜", "红萝卜"}
        else:
            continue
        if not valid:
            greenhouse_violations.append(key)
    checks.append(
        _check(
            "greenhouse_rules",
            float(len(greenhouse_violations)),
            0.0,
            f"illegal greenhouse cells={greenhouse_violations[:10]}",
        )
    )

    regime_violation = 0.0
    simultaneous_regimes: list[tuple[str, int]] = []
    illegal_second_choice: list[tuple[str, int, list[int]]] = []
    for plot in data.irrigated_plots:
        area = data.plots[plot].area_mu
        for year in YEARS:
            u = regimes[(plot, year)]
            rice_area = area_by_plot_season[(plot, year, SEASON_SINGLE)]
            first_area = area_by_plot_season[(plot, year, SEASON_FIRST)]
            second_area = area_by_plot_season[(plot, year, SEASON_SECOND)]
            regime_violation = max(
                regime_violation,
                rice_area - area * u,
                first_area - area * (1.0 - u),
                second_area - area * (1.0 - u),
            )
            if rice_area > tolerance and (first_area > tolerance or second_area > tolerance):
                simultaneous_regimes.append((plot, year))
            second_crops = [
                crop_id
                for crop_id in data.compatible_crops(plot, SEASON_SECOND)
                if x[(plot, crop_id, year, SEASON_SECOND)] > tolerance
            ]
            if len(second_crops) > 1:
                illegal_second_choice.append((plot, year, second_crops))
    checks.append(
        _check(
            "irrigated_annual_regime",
            max(regime_violation, float(len(simultaneous_regimes))),
            tolerance,
            f"simultaneous regime uses={simultaneous_regimes[:10]}",
        )
    )
    checks.append(
        _check(
            "irrigated_second_season_single_crop",
            float(len(illegal_second_choice)),
            0.0,
            f"violations={illegal_second_choice[:10]}",
        )
    )

    rotation_violations: list[tuple[Any, int, float, float]] = []
    for edge in data.adjacency_edges:
        common = set(data.compatible_crops(edge.plot, edge.season_from)).intersection(
            data.compatible_crops(edge.plot, edge.season_to)
        )
        for crop_id in common:
            first = x[(edge.plot, crop_id, edge.year_from, edge.season_from)]
            second = x[(edge.plot, crop_id, edge.year_to, edge.season_to)]
            if first > tolerance and second > tolerance:
                rotation_violations.append((edge, crop_id, first, second))
    checks.append(
        _check(
            "adjacent_opportunity_rotation",
            float(len(rotation_violations)),
            0.0,
            f"violations={rotation_violations[:5]}",
        )
    )

    initial_violations: list[PlantKey] = []
    for plot, previous_crops in data.initial_last_crops.items():
        for season in data.initial_first_seasons[plot]:
            allowed = set(data.compatible_crops(plot, season))
            for crop_id in previous_crops.intersection(allowed):
                key = (plot, crop_id, 2024, season)
                if x[key] > tolerance:
                    initial_violations.append(key)
    checks.append(
        _check(
            "rotation_2023_to_2024",
            float(len(initial_violations)),
            0.0,
            f"violations={initial_violations[:10]}",
        )
    )

    bean_violations: list[tuple[str, str, float, float]] = []
    bean_ids = data.bean_crop_ids
    for plot, plot_data in data.plots.items():
        first_total = data.bean_area_2023[plot] + math.fsum(
            area
            for (cell_plot, crop_id, year, _season), area in x.items()
            if cell_plot == plot and crop_id in bean_ids and year in {2024, 2025}
        )
        if first_total + tolerance < plot_data.area_mu:
            bean_violations.append(
                (plot, "2023-2025", first_total, plot_data.area_mu)
            )
        for start in range(2024, 2029):
            total = math.fsum(
                area
                for (cell_plot, crop_id, year, _season), area in x.items()
                if cell_plot == plot
                and crop_id in bean_ids
                and start <= year <= start + 2
            )
            if total + tolerance < plot_data.area_mu:
                bean_violations.append(
                    (plot, f"{start}-{start + 2}", total, plot_data.area_mu)
                )
    bean_max_violation = max(
        (required - actual for _plot, _window, actual, required in bean_violations),
        default=0.0,
    )
    checks.append(
        _check(
            "rolling_three_year_bean_area",
            bean_max_violation,
            tolerance,
            f"violating windows={bean_violations[:10]}",
        )
    )

    production = _production_by_sales_group(data, x)
    balance_violation = max(
        (
            abs(production[key] - normal[key] - excess[key])
            for key in production
        ),
        default=0.0,
    )
    checks.append(
        _check(
            "production_sales_balance",
            balance_violation,
            tolerance,
            f"checked {len(production)} crop-land-year-season groups",
        )
    )

    total_production_by_demand_key: dict[tuple[int, int, str], float] = defaultdict(float)
    total_normal_by_demand_key: dict[tuple[int, int, str], float] = defaultdict(float)
    for (crop_id, _land, year, season), value in production.items():
        total_production_by_demand_key[(crop_id, year, season)] += value
    for (crop_id, _land, year, season), value in normal.items():
        total_normal_by_demand_key[(crop_id, year, season)] += value
    sales_limit_violation = 0.0
    normal_fill_violation = 0.0
    for key, produced in total_production_by_demand_key.items():
        crop_id, _year, season = key
        demand = data.expected_sales_proxy.get((crop_id, season), 0.0)
        sold_normal = total_normal_by_demand_key[key]
        sales_limit_violation = max(sales_limit_violation, sold_normal - demand)
        normal_fill_violation = max(
            normal_fill_violation, abs(sold_normal - min(produced, demand))
        )
    checks.append(
        _check(
            "expected_sales_cap",
            sales_limit_violation,
            tolerance,
            "normal sales aggregated across land types do not exceed crop-season proxy",
        )
    )
    checks.append(
        _check(
            "normal_sales_filled_before_excess",
            normal_fill_violation,
            tolerance,
            "normal sales equal min(total production, expected-sales proxy)",
        )
    )

    revenue = 0.0
    for key, normal_value in normal.items():
        crop_id, land_type, _year, season = key
        price = data.parameters[(crop_id, land_type, season)].midpoint_price_yuan_per_jin
        revenue += price * normal_value
        revenue += price * solution.model.excess_price_factor * excess[key]
    planting_cost = math.fsum(
        area
        * data.parameters[(crop_id, data.plots[plot].land_type, season)].cost_yuan_per_mu
        for (plot, crop_id, _year, season), area in x.items()
    )
    recomputed_profit = revenue - planting_cost
    objective_violation = abs(recomputed_profit - solution.total_profit_yuan)
    checks.append(
        _check(
            "profit_reconciliation",
            objective_violation,
            max(tolerance, abs(solution.total_profit_yuan) * 1e-9),
            f"recomputed={recomputed_profit:.8f}, solver={solution.total_profit_yuan:.8f}",
        )
    )

    production_total = math.fsum(production.values())
    normal_total = math.fsum(normal.values())
    excess_total = math.fsum(excess.values())
    all_passed = all(check.passed for check in checks)
    return ValidationReport(
        case_name=solution.model.case_name,
        passed=all_passed,
        tolerance=tolerance,
        checks=tuple(checks),
        recomputed_profit_yuan=recomputed_profit,
        solver_profit_yuan=solution.total_profit_yuan,
        production_total_jin=production_total,
        normal_sales_total_jin=normal_total,
        excess_total_jin=excess_total,
    )


def require_valid(report: ValidationReport) -> None:
    if report.passed:
        return
    failures = [
        f"{check.name}: violation={check.max_violation}, {check.details}"
        for check in report.checks
        if not check.passed
    ]
    raise Q1ValidationError(
        f"{report.case_name} failed hard validation: " + "; ".join(failures)
    )


def expected_template_areas(
    data: Q1Data, x: dict[PlantKey, float], year: int
) -> tuple[dict[tuple[str, int], float], dict[tuple[str, int], float]]:
    first: dict[tuple[str, int], float] = defaultdict(float)
    second: dict[tuple[str, int], float] = defaultdict(float)
    for (plot, crop_id, cell_year, season), area in x.items():
        if cell_year != year:
            continue
        if season in {SEASON_SINGLE, SEASON_FIRST}:
            first[(plot, crop_id)] += area
        elif season == SEASON_SECOND:
            second[(plot, crop_id)] += area
    return dict(first), dict(second)


def validate_filled_workbook(
    workbook_path: Path,
    template_path: Path,
    data: Q1Data,
    x: dict[PlantKey, float],
    *,
    tolerance: float = 1e-5,
) -> dict[str, Any]:
    if not workbook_path.is_file():
        raise Q1ValidationError(f"Filled workbook not found: {workbook_path}")
    original_hash_before = sha256(template_path)
    package_validation = validate_template_package(workbook_path, template_path)
    # Random cell access on openpyxl's read-only streaming worksheets reparses
    # rows repeatedly and is quadratic here. Normal in-memory loading remains
    # non-mutating because these workbooks are never saved by the validator.
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    template = load_workbook(template_path, data_only=True, read_only=False)
    plot_order = list(data.plots)
    second_plot_order = [
        plot for plot in plot_order if data.plots[plot].land_type not in OPEN_LAND_TYPES
    ]
    max_difference = 0.0
    checked_cells = 0
    structure_mismatches: list[str] = []
    try:
        if workbook.sheetnames != [str(year) for year in YEARS]:
            structure_mismatches.append(f"sheet names={workbook.sheetnames}")
        for year in YEARS:
            sheet = workbook[str(year)]
            source_sheet = template[str(year)]
            first, second = expected_template_areas(data, x, year)
            for column, crop_id in enumerate(sorted(data.crops), start=3):
                if normalize_cell_text(sheet.cell(1, column).value) != data.crops[crop_id].name:
                    structure_mismatches.append(
                        f"{year}!{sheet.cell(1, column).coordinate} crop header mismatch"
                    )
                for row, plot in enumerate(plot_order, start=2):
                    if normalize_cell_text(sheet.cell(row, 2).value) != plot:
                        structure_mismatches.append(
                            f"{year}!B{row} plot mismatch"
                        )
                    actual = numeric_cell(sheet.cell(row, column).value)
                    expected = first.get((plot, crop_id), 0.0)
                    max_difference = max(max_difference, abs(actual - expected))
                    checked_cells += 1
                for row, plot in enumerate(second_plot_order, start=56):
                    if normalize_cell_text(sheet.cell(row, 2).value) != plot:
                        structure_mismatches.append(
                            f"{year}!B{row} second plot mismatch"
                        )
                    actual = numeric_cell(sheet.cell(row, column).value)
                    expected = second.get((plot, crop_id), 0.0)
                    max_difference = max(max_difference, abs(actual - expected))
                    checked_cells += 1
            for row in range(1, 88):
                for column in (1, 2):
                    if sheet.cell(row, column).value != source_sheet.cell(row, column).value:
                        structure_mismatches.append(
                            f"{year}!{sheet.cell(row, column).coordinate} label changed"
                        )
    finally:
        workbook.close()
        template.close()
    original_hash_after = sha256(template_path)
    if original_hash_before != original_hash_after:
        structure_mismatches.append("original template hash changed during validation")
    passed = (
        max_difference <= tolerance
        and not structure_mismatches
        and package_validation["passed"]
    )
    report = {
        "passed": passed,
        "checked_numeric_cells": checked_cells,
        "max_abs_difference_mu": max_difference,
        "structure_mismatch_count": len(structure_mismatches),
        "structure_mismatches": structure_mismatches[:50],
        "original_template_sha256": original_hash_after,
        "filled_workbook_sha256": sha256(workbook_path),
        "template_package_validation": package_validation,
    }
    if not passed:
        raise Q1ValidationError(f"Filled workbook validation failed: {report}")
    return report


def validate_template_package(workbook_path: Path, template_path: Path) -> dict[str, Any]:
    """Verify unsupported official-template OOXML survives openpyxl output intact."""
    sheet_pattern = re.compile(r"xl/worksheets/sheet\d+\.xml")
    sheet_data_pattern = re.compile(rb"<sheetData(?:\s[^>]*)?>.*?</sheetData>", re.DOTALL)
    mismatches: list[str] = []
    with zipfile.ZipFile(template_path, "r") as template_zip, zipfile.ZipFile(
        workbook_path, "r"
    ) as workbook_zip:
        template_names = template_zip.namelist()
        workbook_names = workbook_zip.namelist()
        if template_names != workbook_names:
            mismatches.append("OOXML package part order/names differ from official template")
        for name in template_names:
            if name not in workbook_names or name == "xl/styles.xml":
                continue
            template_content = template_zip.read(name)
            workbook_content = workbook_zip.read(name)
            if sheet_pattern.fullmatch(name):
                template_content, template_count = sheet_data_pattern.subn(
                    b"<sheetData/>", template_content, count=1
                )
                workbook_content, workbook_count = sheet_data_pattern.subn(
                    b"<sheetData/>", workbook_content, count=1
                )
                if template_count != 1 or workbook_count != 1:
                    mismatches.append(f"{name}: expected exactly one sheetData element")
                    continue
            if template_content != workbook_content:
                mismatches.append(f"{name}: non-cell template content changed")
    return {
        "passed": not mismatches,
        "preservation_rule": (
            "all original OOXML parts byte-identical except xl/styles.xml and "
            "worksheet sheetData elements"
        ),
        "mismatch_count": len(mismatches),
        "mismatches": mismatches[:50],
    }


def normalize_cell_text(value: object) -> str:
    if value is None:
        return ""
    return "".join(str(value).split())


def numeric_cell(value: object) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise Q1ValidationError(f"Expected numeric/blank Excel cell, got {value!r}") from exc
