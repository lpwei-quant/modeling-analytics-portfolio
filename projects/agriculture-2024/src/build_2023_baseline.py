#!/usr/bin/env python3
"""Reconstruct the 2023 crop-season baseline from the local CUMCM files.

The script never modifies the source workbooks.  It reads the two attachments,
maps every 2023 planting record to exactly one parameter row, calculates
record-level production/cost/midpoint revenue, and writes a crop-season
aggregate used as the provisional expected-sales baseline for Question 1.

Modeling assumptions implemented here:
1. 2023 theoretical production proxies 2023 expected sales.
2. The midpoint of each observed price interval is the representative price.

The smart-greenhouse first-season parameter rows are absent from Attachment 2.
The workbook note explicitly states that those parameters equal the ordinary-
greenhouse first-season values, so those values are derived through that stated
mapping rather than treated as directly listed rows.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT_1 = PROJECT_ROOT / "data" / "附件1 (1).xlsx"
ATTACHMENT_2 = PROJECT_ROOT / "data" / "附件2 (2).xlsx"
DEFAULT_OUTPUT = PROJECT_ROOT / "research" / "baseline_2023.csv"

LAND_SHEET = "乡村的现有耕地"
CROP_SHEET = "乡村种植的农作物"
PLANTING_SHEET = "2023年的农作物种植情况"
PARAMETER_SHEET = "2023年统计的相关数据"

GREENHOUSE_TYPES = {"普通大棚", "智慧大棚"}
EXPECTED_PLOT_COUNT = 54
EXPECTED_CROP_COUNT = 41
EXPECTED_TOTAL_AREA_MU = 1213.0
EXPECTED_OPEN_AREA_MU = 1201.0
EXPECTED_GREENHOUSE_AREA_MU = 12.0
NUMERIC_TOLERANCE = 1e-9


class BaselineError(RuntimeError):
    """Raised when the source data fail a required reconstruction invariant."""


def clean_text(value: Any) -> str:
    """Normalize visible whitespace in identifiers and categorical labels."""
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value))


def require_number(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise BaselineError(f"{label} must be numeric, got {value!r}")
    number = float(value)
    if not math.isfinite(number):
        raise BaselineError(f"{label} must be finite, got {value!r}")
    return number


def require_int(value: Any, label: str) -> int:
    number = require_number(value, label)
    integer = int(number)
    if abs(number - integer) > NUMERIC_TOLERANCE:
        raise BaselineError(f"{label} must be an integer, got {value!r}")
    return integer


def parse_price_interval(value: Any, label: str) -> tuple[float, float]:
    """Parse an observed 'lower-upper' price interval without choosing a side."""
    text = str(value).strip() if value is not None else ""
    values = re.findall(r"\d+(?:\.\d+)?", text)
    if len(values) != 2:
        raise BaselineError(f"{label} must contain exactly two bounds, got {value!r}")
    lower, upper = map(float, values)
    if lower < 0 or upper < 0 or lower > upper:
        raise BaselineError(f"{label} has an invalid non-negative interval: {value!r}")
    return lower, upper


@dataclass(frozen=True)
class Plot:
    name: str
    land_type: str
    area_mu: float


@dataclass(frozen=True)
class Crop:
    crop_id: int
    name: str
    crop_type: str


@dataclass(frozen=True)
class Parameter:
    source_row: int
    crop_id: int
    crop_name: str
    land_type: str
    season: str
    yield_jin_per_mu: float
    cost_yuan_per_mu: float
    price_lower_yuan_per_jin: float
    price_upper_yuan_per_jin: float

    @property
    def midpoint_price(self) -> float:
        return (self.price_lower_yuan_per_jin + self.price_upper_yuan_per_jin) / 2.0


@dataclass(frozen=True)
class MappedRecord:
    source_row: int
    plot: Plot
    crop: Crop
    season: str
    planted_area_mu: float
    parameter: Parameter
    parameter_land_type: str
    mapping_rule: str

    @property
    def production_jin(self) -> float:
        return self.planted_area_mu * self.parameter.yield_jin_per_mu

    @property
    def planting_cost_yuan(self) -> float:
        return self.planted_area_mu * self.parameter.cost_yuan_per_mu

    @property
    def midpoint_revenue_yuan(self) -> float:
        return self.production_jin * self.parameter.midpoint_price


def open_data_only(path: Path):
    if not path.is_file():
        raise BaselineError(f"Required source workbook not found: {path}")
    return load_workbook(path, data_only=True, read_only=False)


def read_plots(workbook) -> dict[str, Plot]:
    ws = workbook[LAND_SHEET]
    plots: dict[str, Plot] = {}
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_name, raw_land_type, raw_area = row[:3]
        name = clean_text(raw_name)
        if not name:
            continue
        land_type = clean_text(raw_land_type)
        area_mu = require_number(raw_area, f"{LAND_SHEET}!C{row_number}")
        if area_mu < 0:
            raise BaselineError(f"Negative plot area at {LAND_SHEET}!C{row_number}")
        if name in plots:
            raise BaselineError(f"Duplicate plot name {name!r} in {LAND_SHEET}")
        plots[name] = Plot(name=name, land_type=land_type, area_mu=area_mu)
    return plots


def read_crops(workbook) -> dict[int, Crop]:
    ws = workbook[CROP_SHEET]
    crops: dict[int, Crop] = {}
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_id, raw_name, raw_type = row[:3]
        if raw_id is None:
            continue
        if clean_text(raw_id).startswith("注"):
            continue
        crop_id = require_int(raw_id, f"{CROP_SHEET}!A{row_number}")
        name = clean_text(raw_name)
        crop_type = clean_text(raw_type)
        if not name:
            raise BaselineError(f"Missing crop name at {CROP_SHEET}!B{row_number}")
        if crop_id in crops:
            raise BaselineError(f"Duplicate crop id {crop_id} in {CROP_SHEET}")
        crops[crop_id] = Crop(crop_id=crop_id, name=name, crop_type=crop_type)
    return crops


def read_parameters(workbook) -> tuple[dict[tuple[int, str, str], list[Parameter]], list[str]]:
    ws = workbook[PARAMETER_SHEET]
    parameters: dict[tuple[int, str, str], list[Parameter]] = defaultdict(list)
    parse_issues: list[str] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_sequence, raw_crop_id, raw_crop_name, raw_land_type, raw_season = row[:5]
        if raw_sequence is None:
            continue
        if clean_text(raw_sequence).startswith("注"):
            continue
        try:
            require_int(raw_sequence, f"{PARAMETER_SHEET}!A{row_number}")
            crop_id = require_int(raw_crop_id, f"{PARAMETER_SHEET}!B{row_number}")
            crop_name = clean_text(raw_crop_name)
            land_type = clean_text(raw_land_type)
            season = clean_text(raw_season)
            yield_value = require_number(row[5], f"{PARAMETER_SHEET}!F{row_number}")
            cost_value = require_number(row[6], f"{PARAMETER_SHEET}!G{row_number}")
            price_lower, price_upper = parse_price_interval(
                row[7], f"{PARAMETER_SHEET}!H{row_number}"
            )
            if yield_value < 0 or cost_value < 0:
                raise BaselineError(
                    f"Negative yield or cost in {PARAMETER_SHEET} row {row_number}"
                )
            parameter = Parameter(
                source_row=row_number,
                crop_id=crop_id,
                crop_name=crop_name,
                land_type=land_type,
                season=season,
                yield_jin_per_mu=yield_value,
                cost_yuan_per_mu=cost_value,
                price_lower_yuan_per_jin=price_lower,
                price_upper_yuan_per_jin=price_upper,
            )
            parameters[(crop_id, land_type, season)].append(parameter)
        except BaselineError as exc:
            parse_issues.append(str(exc))
    return dict(parameters), parse_issues


def parameter_key_for(plot_land_type: str, season: str, crop_id: int) -> tuple[tuple[int, str, str], str]:
    if plot_land_type == "智慧大棚" and season == "第一季":
        return (crop_id, "普通大棚", season), "附件2注释：智慧大棚第一季参数等同普通大棚第一季"
    return (crop_id, plot_land_type, season), "直接按作物×地类×季次连接"


def read_and_map_plantings(
    workbook,
    plots: dict[str, Plot],
    crops: dict[int, Crop],
    parameters: dict[tuple[int, str, str], list[Parameter]],
) -> tuple[list[MappedRecord], list[str], list[str]]:
    ws = workbook[PLANTING_SHEET]
    records: list[MappedRecord] = []
    unmatched: list[str] = []
    duplicated: list[str] = []
    current_plot_name = ""

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw_plot, raw_crop_id, raw_crop_name, _raw_crop_type, raw_area, raw_season = row[:6]
        if all(value is None for value in row[:6]):
            continue

        # Excel stores only the first cell of each merged plot-name range.
        if clean_text(raw_plot):
            current_plot_name = clean_text(raw_plot)
        if not current_plot_name:
            unmatched.append(f"row {row_number}: no plot name available after forward-fill")
            continue

        try:
            crop_id = require_int(raw_crop_id, f"{PLANTING_SHEET}!B{row_number}")
            crop_name = clean_text(raw_crop_name)
            season = clean_text(raw_season)
            planted_area_mu = require_number(raw_area, f"{PLANTING_SHEET}!E{row_number}")
            if planted_area_mu < 0:
                raise BaselineError(f"negative planted area in row {row_number}")
        except BaselineError as exc:
            unmatched.append(f"row {row_number}: {exc}")
            continue

        plot = plots.get(current_plot_name)
        crop = crops.get(crop_id)
        if plot is None:
            unmatched.append(f"row {row_number}: unknown plot {current_plot_name!r}")
            continue
        if crop is None:
            unmatched.append(f"row {row_number}: unknown crop id {crop_id}")
            continue
        if crop.name != crop_name:
            unmatched.append(
                f"row {row_number}: crop name {crop_name!r} does not match master {crop.name!r}"
            )
            continue

        parameter_key, mapping_rule = parameter_key_for(plot.land_type, season, crop_id)
        candidates = parameters.get(parameter_key, [])
        if len(candidates) == 0:
            unmatched.append(
                f"row {row_number}: no parameter for crop={crop_id}, "
                f"plot_land_type={plot.land_type}, season={season}, lookup={parameter_key}"
            )
            continue
        if len(candidates) > 1:
            duplicated.append(
                f"row {row_number}: {len(candidates)} parameter rows for lookup={parameter_key}: "
                + ",".join(str(item.source_row) for item in candidates)
            )
            continue

        parameter = candidates[0]
        if parameter.crop_name != crop.name:
            unmatched.append(
                f"row {row_number}: parameter crop name {parameter.crop_name!r} "
                f"does not match crop master {crop.name!r}"
            )
            continue

        records.append(
            MappedRecord(
                source_row=row_number,
                plot=plot,
                crop=crop,
                season=season,
                planted_area_mu=planted_area_mu,
                parameter=parameter,
                parameter_land_type=parameter.land_type,
                mapping_rule=mapping_rule,
            )
        )
    return records, unmatched, duplicated


def validate_source_dimensions(plots: dict[str, Plot], crops: dict[int, Crop]) -> dict[str, Any]:
    total_area = math.fsum(plot.area_mu for plot in plots.values())
    greenhouse_area = math.fsum(
        plot.area_mu for plot in plots.values() if plot.land_type in GREENHOUSE_TYPES
    )
    open_area = math.fsum(
        plot.area_mu for plot in plots.values() if plot.land_type not in GREENHOUSE_TYPES
    )
    checks = {
        "plot_count": len(plots),
        "crop_count": len(crops),
        "total_area_mu": total_area,
        "open_area_mu": open_area,
        "greenhouse_area_mu": greenhouse_area,
    }
    expected = {
        "plot_count": EXPECTED_PLOT_COUNT,
        "crop_count": EXPECTED_CROP_COUNT,
        "total_area_mu": EXPECTED_TOTAL_AREA_MU,
        "open_area_mu": EXPECTED_OPEN_AREA_MU,
        "greenhouse_area_mu": EXPECTED_GREENHOUSE_AREA_MU,
    }
    failures = []
    for name, observed in checks.items():
        target = expected[name]
        if isinstance(target, int):
            passed = observed == target
        else:
            passed = math.isclose(observed, target, rel_tol=0.0, abs_tol=NUMERIC_TOLERANCE)
        if not passed:
            failures.append(f"{name}: observed={observed}, expected={target}")
    if failures:
        raise BaselineError("Source dimension checks failed: " + "; ".join(failures))
    return checks


def find_capacity_violations(records: Iterable[MappedRecord]) -> list[dict[str, Any]]:
    used_area: dict[tuple[str, str], float] = defaultdict(float)
    plot_lookup: dict[str, Plot] = {}
    for record in records:
        used_area[(record.plot.name, record.season)] += record.planted_area_mu
        plot_lookup[record.plot.name] = record.plot
    violations = []
    for (plot_name, season), used in sorted(used_area.items()):
        capacity = plot_lookup[plot_name].area_mu
        if used > capacity + NUMERIC_TOLERANCE:
            violations.append(
                {
                    "plot": plot_name,
                    "season": season,
                    "used_area_mu": used,
                    "capacity_mu": capacity,
                    "excess_mu": used - capacity,
                }
            )
    return violations


def aggregate_crop_season(records: Iterable[MappedRecord]) -> list[dict[str, Any]]:
    grouped: dict[tuple[int, str, str], dict[str, Any]] = {}
    for record in records:
        key = (record.crop.crop_id, record.crop.name, record.season)
        group = grouped.setdefault(
            key,
            {
                "crop_id": record.crop.crop_id,
                "crop_name": record.crop.name,
                "planting_season": record.season,
                "source_record_count": 0,
                "total_planted_area_mu": 0.0,
                "theoretical_production_jin": 0.0,
                "baseline_expected_sales_proxy_jin": 0.0,
                "total_planting_cost_yuan": 0.0,
                "midpoint_revenue_yuan": 0.0,
                "smart_greenhouse_first_season_derived_record_count": 0,
            },
        )
        group["source_record_count"] += 1
        group["total_planted_area_mu"] += record.planted_area_mu
        group["theoretical_production_jin"] += record.production_jin
        group["baseline_expected_sales_proxy_jin"] += record.production_jin
        group["total_planting_cost_yuan"] += record.planting_cost_yuan
        group["midpoint_revenue_yuan"] += record.midpoint_revenue_yuan
        if record.plot.land_type == "智慧大棚" and record.season == "第一季":
            group["smart_greenhouse_first_season_derived_record_count"] += 1

    season_order = {"单季": 0, "第一季": 1, "第二季": 2}
    output = []
    for group in grouped.values():
        production = group["theoretical_production_jin"]
        group["production_weighted_midpoint_price_yuan_per_jin"] = (
            group["midpoint_revenue_yuan"] / production if production > 0 else 0.0
        )
        group["expected_sales_proxy_assumption"] = "2023理论产量代理2023预期销售量"
        group["price_assumption"] = "观测销售价格区间中点"
        output.append(group)
    output.sort(
        key=lambda row: (
            row["crop_id"], season_order.get(row["planting_season"], 99), row["planting_season"]
        )
    )
    return output


def assert_nonnegative(records: Iterable[MappedRecord], aggregate_rows: Iterable[dict[str, Any]]) -> None:
    for record in records:
        values = {
            "planted_area_mu": record.planted_area_mu,
            "yield_jin_per_mu": record.parameter.yield_jin_per_mu,
            "cost_yuan_per_mu": record.parameter.cost_yuan_per_mu,
            "price_lower": record.parameter.price_lower_yuan_per_jin,
            "price_upper": record.parameter.price_upper_yuan_per_jin,
            "midpoint_price": record.parameter.midpoint_price,
            "production_jin": record.production_jin,
            "planting_cost_yuan": record.planting_cost_yuan,
        }
        negative = {name: value for name, value in values.items() if value < -NUMERIC_TOLERANCE}
        if negative:
            raise BaselineError(f"Negative calculated value in planting row {record.source_row}: {negative}")
    for row in aggregate_rows:
        for name in (
            "total_planted_area_mu",
            "theoretical_production_jin",
            "baseline_expected_sales_proxy_jin",
            "total_planting_cost_yuan",
            "midpoint_revenue_yuan",
            "production_weighted_midpoint_price_yuan_per_jin",
        ):
            if float(row[name]) < -NUMERIC_TOLERANCE:
                raise BaselineError(f"Negative aggregate {name} for {row['crop_name']}/{row['planting_season']}")


def write_csv(rows: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "crop_id",
        "crop_name",
        "planting_season",
        "source_record_count",
        "total_planted_area_mu",
        "theoretical_production_jin",
        "baseline_expected_sales_proxy_jin",
        "total_planting_cost_yuan",
        "production_weighted_midpoint_price_yuan_per_jin",
        "midpoint_revenue_yuan",
        "smart_greenhouse_first_season_derived_record_count",
        "expected_sales_proxy_assumption",
        "price_assumption",
    ]
    with output_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            formatted = dict(row)
            for field in (
                "total_planted_area_mu",
                "theoretical_production_jin",
                "baseline_expected_sales_proxy_jin",
                "total_planting_cost_yuan",
                "production_weighted_midpoint_price_yuan_per_jin",
                "midpoint_revenue_yuan",
            ):
                formatted[field] = f"{float(row[field]):.10f}".rstrip("0").rstrip(".")
            writer.writerow(formatted)


def build_baseline(output_path: Path) -> dict[str, Any]:
    attachment_1 = open_data_only(ATTACHMENT_1)
    attachment_2 = open_data_only(ATTACHMENT_2)
    try:
        plots = read_plots(attachment_1)
        crops = read_crops(attachment_1)
        dimensions = validate_source_dimensions(plots, crops)
        parameters, parameter_parse_issues = read_parameters(attachment_2)
        duplicate_parameter_keys = {
            str(key): [item.source_row for item in values]
            for key, values in parameters.items()
            if len(values) != 1
        }
        records, unmatched, duplicated_joins = read_and_map_plantings(
            attachment_2, plots, crops, parameters
        )
        capacity_violations = find_capacity_violations(records)
        aggregate_rows = aggregate_crop_season(records)
        assert_nonnegative(records, aggregate_rows)

        failures = []
        if parameter_parse_issues:
            failures.append(f"parameter parse issues: {parameter_parse_issues}")
        if duplicate_parameter_keys:
            failures.append(f"duplicate parameter keys: {duplicate_parameter_keys}")
        if unmatched:
            failures.append(f"unmatched planting rows: {unmatched}")
        if duplicated_joins:
            failures.append(f"duplicated joins: {duplicated_joins}")
        if capacity_violations:
            failures.append(f"plot-season capacity violations: {capacity_violations}")
        if failures:
            raise BaselineError("; ".join(failures))

        write_csv(aggregate_rows, output_path)
        return {
            "status": "ok",
            "sources": {
                "attachment_1": str(ATTACHMENT_1.relative_to(PROJECT_ROOT)),
                "attachment_2": str(ATTACHMENT_2.relative_to(PROJECT_ROOT)),
            },
            "output": str(output_path.resolve()),
            **dimensions,
            "parameter_row_count": sum(len(values) for values in parameters.values()),
            "duplicate_parameter_key_count": len(duplicate_parameter_keys),
            "planting_record_count": len(records),
            "mapped_record_count": len(records),
            "unmatched_join_count": len(unmatched),
            "duplicated_join_count": len(duplicated_joins),
            "capacity_violation_count": len(capacity_violations),
            "crop_season_baseline_row_count": len(aggregate_rows),
            "smart_greenhouse_first_season_derived_record_count": sum(
                1
                for record in records
                if record.plot.land_type == "智慧大棚" and record.season == "第一季"
            ),
            "all_calculated_values_nonnegative": True,
            "total_theoretical_production_jin": math.fsum(
                row["theoretical_production_jin"] for row in aggregate_rows
            ),
            "total_planting_cost_yuan": math.fsum(
                row["total_planting_cost_yuan"] for row in aggregate_rows
            ),
        }
    finally:
        attachment_1.close()
        attachment_2.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"CSV output path (default: {DEFAULT_OUTPUT})",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_path = args.output
    if not output_path.is_absolute():
        output_path = PROJECT_ROOT / output_path
    try:
        diagnostics = build_baseline(output_path)
    except BaselineError as exc:
        # ASCII-safe JSON keeps diagnostics legible under Windows consoles with
        # differing active code pages; source/output files remain UTF-8.
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=True, indent=2))
        return 1
    print(json.dumps(diagnostics, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
