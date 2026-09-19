#!/usr/bin/env python3
"""Rebuild, solve, validate, and export both deterministic Q1 cases."""

from __future__ import annotations

import json
import os
import re
import sys
import time
import warnings
import zipfile
from pathlib import Path


CANONICAL_PYTHON = Path(sys.executable)


def require_canonical_python() -> None:
    """Submission copy uses the active Python interpreter."""
    return


require_canonical_python()

# Apply deterministic thread limits before importing NumPy/SciPy.
os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("OPENBLAS_NUM_THREADS", "1")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from build_2023_baseline import build_baseline  # noqa: E402
from q1_data import YEARS, load_q1_data  # noqa: E402
from q1_model import build_q1_model, solve_q1_model  # noqa: E402
from q1_outputs import (  # noqa: E402
    build_metrics,
    build_template_payload,
    solution_long_frame,
    write_case_outputs,
)
from q1_validate import (  # noqa: E402
    require_valid,
    validate_filled_workbook,
    validate_solution,
)


OUTPUT_DIR = PROJECT_ROOT / "outputs"
RESEARCH_DIR = PROJECT_ROOT / "research"
TMP_DIR = PROJECT_ROOT / "tmp" / "q1_delivery"


CASES = (
    {
        "name": "case1_waste",
        "stem": "q1_case1",
        "factor": 0.0,
        "template": "result1_1.xlsx",
        "workbook": "result1_1_filled.xlsx",
    },
    {
        "name": "case2_half_price",
        "stem": "q1_case2",
        "factor": 0.5,
        "template": "result1_2.xlsx",
        "workbook": "result1_2_filled.xlsx",
    },
)


def write_json(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def runtime_diagnostic() -> dict[str, object]:
    import numpy
    import scipy

    return {
        "python_executable": sys.executable,
        "python_version": sys.version.split()[0],
        "numpy_version": numpy.__version__,
        "numpy_path": numpy.__file__,
        "scipy_version": scipy.__version__,
        "scipy_path": scipy.__file__,
        "sys_path": list(sys.path),
    }


def _replace_sheet_data(template_xml: bytes, generated_xml: bytes, part_name: str) -> bytes:
    pattern = re.compile(rb"<sheetData(?:\s[^>]*)?>.*?</sheetData>", re.DOTALL)
    generated_matches = pattern.findall(generated_xml)
    template_matches = pattern.findall(template_xml)
    if len(generated_matches) != 1 or len(template_matches) != 1:
        raise RuntimeError(
            f"{part_name}: expected one sheetData element in template and generated XML"
        )
    return pattern.sub(generated_matches[0], template_xml, count=1)


def preserve_template_package(
    template_path: Path, openpyxl_path: Path, output_path: Path
) -> None:
    """Keep every official-template OOXML part except openpyxl cell/style data."""
    with zipfile.ZipFile(template_path, "r") as template_zip, zipfile.ZipFile(
        openpyxl_path, "r"
    ) as generated_zip, zipfile.ZipFile(output_path, "w") as output_zip:
        generated_names = set(generated_zip.namelist())
        for info in template_zip.infolist():
            content = template_zip.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                if info.filename not in generated_names:
                    raise RuntimeError(f"Generated workbook is missing {info.filename}")
                content = _replace_sheet_data(
                    content, generated_zip.read(info.filename), info.filename
                )
            elif info.filename == "xl/styles.xml":
                content = generated_zip.read(info.filename)
            output_zip.writestr(info, content)


def fill_template(case: dict[str, object], payload_path: Path) -> dict[str, object]:
    from openpyxl import load_workbook
    import openpyxl

    template_path = OUTPUT_DIR / str(case["template"])
    output_path = OUTPUT_DIR / str(case["workbook"])
    openpyxl_path = output_path.with_suffix(".openpyxl.tmp.xlsx")
    temporary_path = output_path.with_suffix(".tmp.xlsx")
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    expected_sheets = [str(year) for year in payload["years"]]
    if len(payload["crop_order"]) != 41:
        raise RuntimeError("Template payload must contain exactly 41 crops")
    if len(payload["plot_order"]) != 54 or len(payload["second_plot_order"]) != 28:
        raise RuntimeError("Template payload must contain 54 first-region and 28 second-region plots")

    with warnings.catch_warnings(record=True) as captured_warnings:
        warnings.simplefilter("always")
        workbook = load_workbook(template_path)
    warning_messages = [str(item.message) for item in captured_warnings]
    unexpected_warnings = [
        message
        for message in warning_messages
        if "Cannot parse header or footer so it will be ignored" not in message
    ]
    if unexpected_warnings:
        workbook.close()
        raise RuntimeError(f"Unexpected openpyxl template warnings: {unexpected_warnings}")
    try:
        if workbook.sheetnames != expected_sheets:
            raise RuntimeError(f"Unexpected worksheets: {workbook.sheetnames}")
        for year in expected_sheets:
            sheet = workbook[year]
            headers = ["".join(str(sheet.cell(1, column).value or "").split()) for column in range(3, 44)]
            first_plots = ["".join(str(sheet.cell(row, 2).value or "").split()) for row in range(2, 56)]
            second_plots = ["".join(str(sheet.cell(row, 2).value or "").split()) for row in range(56, 84)]
            if headers != payload["crop_order"]:
                raise RuntimeError(f"{year}: crop headers changed")
            if first_plots != payload["plot_order"]:
                raise RuntimeError(f"{year}: first-region plot order changed")
            if second_plots != payload["second_plot_order"]:
                raise RuntimeError(f"{year}: second-region plot order changed")

            values = payload["values"][year]
            if (
                len(values["first"]) != 54
                or len(values["second"]) != 28
                or any(len(row) != 41 for row in values["first"] + values["second"])
            ):
                raise RuntimeError(f"{year}: payload value matrix dimensions changed")
            for row_offset, row_values in enumerate(values["first"], start=2):
                for column_offset, value in enumerate(row_values, start=3):
                    cell = sheet.cell(row_offset, column_offset, float(value))
                    cell.number_format = "0.0000"
            for row_offset, row_values in enumerate(values["second"], start=56):
                for column_offset, value in enumerate(row_values, start=3):
                    cell = sheet.cell(row_offset, column_offset, float(value))
                    cell.number_format = "0.0000"
        workbook.save(openpyxl_path)
    finally:
        workbook.close()
    try:
        preserve_template_package(template_path, openpyxl_path, temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        openpyxl_path.unlink(missing_ok=True)
        temporary_path.unlink(missing_ok=True)
    return {
        "status": "ok",
        "engine": "openpyxl",
        "openpyxl_version": openpyxl.__version__,
        "output": str(output_path.resolve()),
        "numeric_cells_written": len(expected_sheets) * (54 + 28) * 41,
        "captured_template_warnings": warning_messages,
        "template_package_preservation": (
            "all original OOXML parts retained; only styles and worksheet sheetData replaced"
        ),
    }


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)

    baseline_diagnostic = build_baseline(RESEARCH_DIR / "baseline_2023.csv")
    data = load_q1_data(PROJECT_ROOT)
    receipt: dict[str, object] = {
        "status": "running",
        "command": f"{sys.executable} scripts/run_q1.py",
        "python_runtime": runtime_diagnostic(),
        "years": list(YEARS),
        "baseline_reconstruction": baseline_diagnostic,
        "cases": [],
    }

    for case in CASES:
        print(f"Building and solving {case['name']}...", flush=True)
        model = build_q1_model(
            data,
            case_name=str(case["name"]),
            excess_price_factor=float(case["factor"]),
        )
        started = time.perf_counter()
        solution = solve_q1_model(
            model,
            time_limit_seconds=300.0,
            mip_relative_gap=0.005,
            display_solver_log=False,
        )
        solve_seconds = time.perf_counter() - started

        validation = validate_solution(solution)
        require_valid(validation)
        frame = solution_long_frame(solution)
        metrics = build_metrics(
            solution, frame, validation, solve_seconds=solve_seconds
        )
        csv_path, metrics_path = write_case_outputs(
            OUTPUT_DIR, str(case["stem"]), frame, metrics
        )

        payload_path = build_template_payload(
            solution, TMP_DIR / f"{case['stem']}_template_payload.json"
        )
        print(f"Writing {case['workbook']} with openpyxl...", flush=True)
        workbook_diagnostic = fill_template(case, payload_path)
        workbook_path = OUTPUT_DIR / str(case["workbook"])
        workbook_validation = validate_filled_workbook(
            workbook_path,
            OUTPUT_DIR / str(case["template"]),
            data,
            solution.x,
        )
        print(f"Workbook cell reconciliation passed for {case['workbook']}.", flush=True)
        metrics["workbook_validation"] = workbook_validation
        metrics["workbook_export"] = workbook_diagnostic
        write_json(metrics_path, metrics)
        receipt["cases"].append(
            {
                "case": case["name"],
                "profit_yuan": metrics["totals"]["profit_yuan"],
                "validation_passed": validation.passed,
                "csv": str(csv_path.relative_to(PROJECT_ROOT)),
                "metrics": str(metrics_path.relative_to(PROJECT_ROOT)),
                "workbook": str(workbook_path.relative_to(PROJECT_ROOT)),
                "workbook_validation_passed": workbook_validation["passed"],
            }
        )
        print(
            f"Completed {case['name']}: profit={solution.total_profit_yuan:.2f}, "
            f"gap={solution.result.mip_gap:.6f}, solve_seconds={solve_seconds:.1f}",
            flush=True,
        )

    receipt["status"] = "ok"
    write_json(OUTPUT_DIR / "q1_run_receipt.json", receipt)
    print("Both Q1 cases solved, validated, and exported.", flush=True)
    return 0


if __name__ == "__main__":
    if sys.argv[1:] == ["--runtime-check"]:
        print(json.dumps(runtime_diagnostic(), ensure_ascii=False, indent=2))
        raise SystemExit(0)
    if sys.argv[1:]:
        raise SystemExit(f"Unsupported arguments: {sys.argv[1:]}")
    raise SystemExit(main())
