"""Exact Q1 LP using storage increments, with an optional lexicographic tie rule."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
from datetime import datetime, time, timezone
from pathlib import Path
from time import perf_counter

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import csr_matrix, hstack, lil_matrix, vstack

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CASES = ["base", "roundtrip90", "left_endpoint", "interval_average",
                 "free_initial", "ideal_efficiency", "power_half", "power_double",
                 "usable_half", "usable_double"]

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def label(minutes):
    return "24:00" if minutes == 1440 else f"{minutes//60:02d}:{minutes%60:02d}"

def read_inputs(config):
    source = ROOT / config["source_workbook"]
    if digest(source) != config["source_sha256"]:
        raise ValueError("Raw workbook hash differs from the audited source")
    workbook = load_workbook(source, data_only=False, read_only=True, keep_links=False)
    rows = list(workbook[config["source_sheet"]].iter_rows(values_only=True))
    workbook.close()
    if rows[0] != ("时间", "电价", "小区负载", "光伏发电预测功率") or len(rows) != 145:
        raise ValueError("Unexpected source schema or row count")
    raw_labels = [r[0].strftime("%H:%M") if isinstance(r[0], time) else r[0] for r in rows[1:]]
    if raw_labels != [label(m) for m in range(10, 1440, 10)] + ["0:00+1"]:
        raise ValueError("Unexpected source timestamps")
    values = np.asarray([r[1:] for r in rows[1:]], dtype=float)
    if not np.isfinite(values).all():
        raise ValueError("Source includes nonfinite values")
    return values, raw_labels

def interpreted(values, mode):
    if mode == "right_endpoint_representative":
        return values.copy()
    if mode == "left_endpoint_periodic":
        return np.roll(values, 1, axis=0)
    if mode == "trapezoid_periodic":
        # A discrete interval-average approximation, not the exact integral of
        # a continuously varying price multiplied by continuously varying power.
        return (values + np.roll(values, 1, axis=0))/2
    raise ValueError(f"Unknown interpretation: {mode}")

def case_config(base, case):
    cfg = dict(base)
    if case == "base":
        pass
    elif case == "roundtrip90":
        cfg.update(eta_charge=math.sqrt(.9), eta_discharge=math.sqrt(.9))
    elif case == "left_endpoint":
        cfg["time_interpretation"] = "left_endpoint_periodic"
    elif case == "interval_average":
        cfg["time_interpretation"] = "trapezoid_periodic"
    elif case == "free_initial":
        cfg["initial_soc_kwh"] = None
    elif case == "ideal_efficiency":
        cfg.update(eta_charge=1., eta_discharge=1.)
    elif case == "power_half":
        cfg.update(max_charge_power_kw=2500., max_discharge_power_kw=2500.)
    elif case == "power_double":
        cfg.update(max_charge_power_kw=10000., max_discharge_power_kw=10000.)
    elif case == "usable_half":
        cfg.update(soc_min_kwh=3600., soc_max_kwh=8400.)
    elif case == "usable_double":
        cfg.update(capacity_kwh=24000., soc_min_kwh=2400., soc_max_kwh=21600., initial_soc_kwh=12000.)
    elif case == "upper_plus_10":
        cfg["soc_max_kwh"] += 10.
    elif case == "lower_minus_10":
        cfg["soc_min_kwh"] -= 10.
    elif case == "charge_power_plus_100":
        cfg["max_charge_power_kw"] += 100.
    elif case == "discharge_power_plus_100":
        cfg["max_discharge_power_kw"] += 100.
    else:
        raise ValueError(f"Unknown case: {case}")
    cfg["case"] = case
    return cfg

def formulate(price, load, pv, config):
    """All energies in kWh; ordering x=[g_1..g_n,S_0..S_n]."""
    price, load, pv = [np.asarray(a, dtype=float) for a in (price, load, pv)]
    n = len(price)
    if not n or price.ndim != 1 or load.shape != price.shape or pv.shape != price.shape:
        raise ValueError("Expected nonempty, equal-length one-dimensional arrays")
    if not all(np.isfinite(a).all() for a in (price, load, pv)):
        raise ValueError("Nonfinite model input")
    if (price <= 0).any() or (load < 0).any() or (pv < 0).any():
        raise ValueError("Exact epigraph model requires positive prices and nonnegative load/PV")
    ec, ed = config["eta_charge"], config["eta_discharge"]
    if not (0 < ec <= 1 and 0 < ed <= 1):
        raise ValueError("Each efficiency must lie in (0,1]")
    dt = config["step_hours"]
    low, high, initial = config["soc_min_kwh"], config["soc_max_kwh"], config["initial_soc_kwh"]
    if dt <= 0 or low > high or (initial is not None and not low <= initial <= high):
        raise ValueError("Invalid time step, state bounds, or initial state")
    cmax, dmax = config["max_charge_power_kw"]*dt, config["max_discharge_power_kw"]*dt
    if cmax < 0 or dmax < 0:
        raise ValueError("Power limits must be nonnegative")
    if (config.get("grid_export_allowed", False) or config.get("battery_dump_allowed", False)
            or not config.get("pv_curtailment_allowed", True)):
        raise ValueError("This formulation implements free PV curtailment, no export and no battery dumping")
    size = 2*n+1
    objective = np.r_[price, np.zeros(n+1)]
    lower = np.r_[np.zeros(n), np.full(n+1, low)]
    upper = np.r_[np.full(n, np.inf), np.full(n+1, high)]
    if initial is not None:
        lower[n] = upper[n] = lower[-1] = upper[-1] = initial
    aub = lil_matrix((4*n, size), dtype=float)
    bub = np.zeros(4*n)
    for t in range(n):
        for row, slope in ((t, 1/ec), (n+t, ed)):
            aub[row, t] = -1
            aub[row, n+t] = -slope
            aub[row, n+t+1] = slope
            bub[row] = pv[t]-load[t]
        aub[2*n+t, n+t] = -1
        aub[2*n+t, n+t+1] = 1
        bub[2*n+t] = ec*cmax
        aub[3*n+t, n+t] = 1
        aub[3*n+t, n+t+1] = -1
        # Under exclusivity, no export and no battery dumping imply d_t<=load_t.
        bub[3*n+t] = min(dmax, load[t])/ed
    aeq = lil_matrix((1, size), dtype=float)
    aeq[0, n] = -1; aeq[0, -1] = 1
    return dict(n=n, c=objective, lower=lower, upper=upper,
                aub=aub.tocsr(), bub=bub, aeq=aeq.tocsr(), beq=np.zeros(1))

def run_lp(model, method="highs-ds"):
    start = perf_counter()
    result = linprog(model["c"], A_ub=model["aub"], b_ub=model["bub"],
                     A_eq=model["aeq"], b_eq=model["beq"],
                     bounds=list(zip(model["lower"], model["upper"])), method=method,
                     options={"primal_feasibility_tolerance": 1e-9,
                              "dual_feasibility_tolerance": 1e-9,
                              "ipm_optimality_tolerance": 1e-10})
    seconds = perf_counter()-start
    if not result.success:
        raise RuntimeError(f"LP failed ({result.status}): {result.message}")
    return result, seconds

def optimal_face(model, optimum, tolerance=1e-9):
    """Use complementary slackness, avoiding an equality to a rounded optimum."""
    result = dict(model)
    active = np.flatnonzero(optimum.ineqlin.marginals < -tolerance)
    result["aeq"] = vstack([model["aeq"],model["aub"][active]]).tocsr()
    result["beq"] = np.r_[model["beq"],model["bub"][active]]
    result["lower"] = model["lower"].copy()
    result["upper"] = model["upper"].copy()
    at_lower = np.flatnonzero(optimum.lower.marginals > tolerance)
    at_upper = np.flatnonzero(optimum.upper.marginals < -tolerance)
    result["upper"][at_lower] = result["lower"][at_lower]
    result["lower"][at_upper] = result["upper"][at_upper]
    result["face_constraints"] = dict(tight_inequality_rows=active.tolist(),
                                      tight_lower_variables=at_lower.tolist(),
                                      tight_upper_variables=at_upper.tolist(),dual_tolerance=tolerance)
    return result

def throughput_model(primary_model, optimum, config):
    """Stage two: minimize exact bus throughput on the primary optimal face."""
    n = primary_model["n"]; old_size = len(primary_model["c"])
    extra = lil_matrix((2*n, old_size+n), dtype=float)
    for t in range(n):
        # h_t >= max(delta/eta_c, -eta_d*delta) = c_t+d_t after recovery.
        for row, slope in ((t, 1/config["eta_charge"]), (n+t, -config["eta_discharge"])):
            extra[row, n+t] = -slope
            extra[row, n+t+1] = slope
            extra[row, old_size+t] = -1
    face = optimal_face(primary_model, optimum)
    aub = vstack([hstack([face["aub"], csr_matrix((4*n, n))]), extra]).tocsr()
    aeq = hstack([face["aeq"],csr_matrix((face["aeq"].shape[0],n))]).tocsr()
    return dict(n=n, c=np.r_[np.zeros(old_size), np.ones(n)],
                lower=np.r_[face["lower"], np.zeros(n)],
                upper=np.r_[face["upper"], np.full(n, np.inf)],
                aub=aub, bub=np.r_[primary_model["bub"], np.zeros(2*n)],
                aeq=aeq, beq=face["beq"],face_constraints=face["face_constraints"])

def recover(x, n, load, pv, config):
    grid = x[:n].copy()
    soc = x[n:2*n+1].copy()
    change = np.diff(soc)
    charge = np.maximum(change, 0)/config["eta_charge"]
    discharge = np.maximum(-change, 0)*config["eta_discharge"]
    spill = grid + pv + discharge - load - charge
    # Remove only solver roundoff; do not change substantive optimized values.
    for array in (grid, charge, discharge, spill):
        array[np.abs(array) < 1e-9] = 0.
    floor = np.maximum(load-pv+charge-discharge, 0.)
    if np.max(np.abs(grid-floor)) > 1e-6:
        raise AssertionError("The grid epigraph is not tight")
    if np.min(spill) < -1e-6 or np.max(spill-pv) > 1e-6:
        raise AssertionError("Recovered unused PV lies outside [0,PV]")
    return dict(g=grid, c=charge, d=discharge, spill=spill, soc=soc, delta=change)

def solve(price, load, pv, config, select=True):
    model = formulate(price, load, pv, config)
    primary, t1 = run_lp(model)
    selected, second_model, t2 = primary, None, 0.
    if select:
        second_model = throughput_model(model, primary, config)
        selected, t2 = run_lp(second_model)
    solution = recover(selected.x, len(price), load, pv, config)
    if abs(float(np.dot(price, solution["g"]))-primary.fun) > 1e-6:
        raise AssertionError("Selection changed the primary objective")
    return dict(model=model, primary=primary, selected=selected, secondary_model=second_model,
                solution=solution, primary_seconds=t1, secondary_seconds=t2)

def metrics(solution, price, load, pv, config):
    total = lambda a: math.fsum(float(v) for v in a)
    g,c,d,s,soc = [solution[k] for k in ("g","c","d","spill","soc")]
    no_battery = np.maximum(load-pv, 0)
    cost, reference = total(g*price), total(no_battery*price)
    return dict(cost_yuan=cost, grid_kwh=total(g), load_kwh=total(load), pv_available_kwh=total(pv),
                pv_curtailed_kwh=total(s), charge_bus_kwh=total(c), discharge_bus_kwh=total(d),
                throughput_kwh=total(c+d), storage_loss_kwh=total((1-config["eta_charge"])*c+(1/config["eta_discharge"]-1)*d),
                initial_soc_kwh=float(soc[0]), terminal_soc_kwh=float(soc[-1]),
                soc_min_kwh=float(min(soc)), soc_max_kwh=float(max(soc)),
                max_charge_power_kw=float(max(c)/config["step_hours"]),
                max_discharge_power_kw=float(max(d)/config["step_hours"]),
                charge_slots=int(sum(c>1e-6)), discharge_slots=int(sum(d>1e-6)),
                simultaneous_charge_discharge_slots=int(sum((c>1e-6)&(d>1e-6))),
                no_battery_grid_kwh=total(no_battery), no_battery_cost_yuan=reference,
                no_battery_curtailed_kwh=total(np.maximum(pv-load, 0)),
                cost_saving_yuan=reference-cost,
                cost_saving_fraction=(reference-cost)/reference if reference else None)
