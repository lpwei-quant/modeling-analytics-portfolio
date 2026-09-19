"""Past-information forecasts and whole-day historical-origin scenarios."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Literal

import numpy as np
from openpyxl import load_workbook

T = 144
DT = 1 / 6
START_DATE = date(2025, 1, 1)

@dataclass(frozen=True)
class DailyData:
    price: np.ndarray
    load: np.ndarray
    pv: np.ndarray
    dates: tuple[date, ...]
    load_template: np.ndarray
    pv_template: np.ndarray

def load_data(root: str | Path | None = None) -> DailyData:
    """Load source workbooks; labels identify interval ends, position is retained."""
    root = Path(root) if root is not None else Path(__file__).resolve().parents[1]
    folder = root / "problem" / "attachments" / "C题"
    book = load_workbook(folder / "附件1.xlsx", read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].iter_rows(min_row=2, values_only=True))
        a1 = np.asarray([r[1:4] for r in rows], dtype=float)
    finally:
        book.close()
    book = load_workbook(folder / "附件2.xlsx", read_only=True, data_only=True)
    try:
        payload = []
        date_sets = []
        for name in ("小区负载", "光伏发电实际功率"):
            rows = list(book[name].iter_rows(min_row=2, values_only=True))
            payload.append(np.asarray([r[1:] for r in rows], dtype=float) * DT)
            date_sets.append(tuple(r[0].date() if isinstance(r[0], datetime) else r[0]
                                   for r in rows))
    finally:
        book.close()
    if a1.shape != (T, 3) or any(x.shape != (365, T) for x in payload):
        raise ValueError("Official source workbook shapes differ from 144/365x144")
    expected_dates = tuple(START_DATE + timedelta(days=j) for j in range(365))
    if not date_sets[0] == date_sets[1] == expected_dates:
        raise ValueError("Source dates are not aligned daily dates in 2025")
    arrays = [a1[:, 0].copy(), *payload, a1[:, 1] * DT, a1[:, 2] * DT]
    for x in arrays:
        if not np.isfinite(x).all() or (x < 0).any():
            raise ValueError("Source contains nonfinite or negative values")
        x.setflags(write=False)
    return DailyData(arrays[0], arrays[1], arrays[2], expected_dates, arrays[3], arrays[4])

def _lowday(index: int) -> bool:
    return (START_DATE + timedelta(days=int(index))).weekday() in (4, 5)

def ar1_coef(residual: np.ndarray) -> float:
    """Pooled adjacent-slot correlation, estimated solely from historical errors."""
    x, y = residual[:, :-1].ravel(), residual[:, 1:].ravel()
    if x.std() < 1e-9 or y.std() < 1e-9:
        return 0.0
    return float(np.clip(np.corrcoef(x, y)[0, 1], 0.0, 0.999))

class ForecastArchive:
    def __init__(self, data: DailyData, cold_start: Literal["reference", "scaled"] = "reference"):
        if cold_start not in ("reference", "scaled"):
            raise ValueError("cold_start must be reference or scaled")
        if data.load.shape != data.pv.shape or data.load.ndim != 2 or data.load.shape[1] != T:
            raise ValueError("load/PV must be aligned day-by-slot arrays")
        self.data = data
        self.cold_start = cold_start
        self._points: dict[tuple[int, int], tuple[np.ndarray, np.ndarray]] = {}
        self._residuals: dict[tuple[int, int], tuple[np.ndarray, np.ndarray, np.ndarray]] = {}
        self._conditional_bases: dict[tuple[int, int, str], tuple] = {}

    def _origin(self, origin: int) -> int:
        if not isinstance(origin, (int, np.integer)) or not 0 <= origin <= len(self.data.load):
            raise ValueError("origin must be a nonnegative index with its past available")
        return int(origin)

    def point(self, origin: int, target: int | None = None) -> tuple[np.ndarray, np.ndarray]:
        """Forecast target day using only actual records strictly before origin."""
        origin = self._origin(origin)
        target = origin if target is None else int(target)
        if target < 0:
            raise ValueError("target must be nonnegative")
        key = (origin, target)
        if key in self._points:
            return self._points[key]
        same = [i for i in range(origin - 1, -1, -1) if _lowday(i) == _lowday(target)][:3]
        if same:
            load_hat = self.data.load[same].mean(axis=0)
        else:
            scale = 78400 / 111025 if _lowday(target) else 124000 / 111025
            if self.cold_start == "reference" and origin == target == 0:
                scale = 1.0
            load_hat = self.data.load_template * scale
        if origin == 0:
            pv_hat = self.data.pv_template.copy()
        else:
            count = min(14, origin)
            weights = np.asarray([0.7 ** j for j in range(count)])
            rows = [origin - 1 - j for j in range(count)]
            pv_hat = (weights[:, None] * self.data.pv[rows]).sum(axis=0) / weights.sum()
        output = (np.maximum(load_hat, 0), np.maximum(pv_hat, 0))
        for x in output:
            x.setflags(write=False)
        self._points[key] = output
        return output

    def residuals(self, origin: int, K: int = 30) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
        """Errors from the forecasts actually available at each historical origin."""
        origin = self._origin(origin)
        if K < 1:
            raise ValueError("K must be positive")
        key = (origin, int(K))
        if key in self._residuals:
            return self._residuals[key]
        indices = np.arange(max(0, origin - K), origin, dtype=int)
        if len(indices):
            e_load, e_pv = [], []
            for j in indices:
                lh, ph = self.point(int(j))
                e_load.append(self.data.load[j] - lh)
                e_pv.append(self.data.pv[j] - ph)
            result = (np.asarray(e_load), np.asarray(e_pv), indices)
        else:
            result = (np.zeros((1, T)), np.zeros((1, T)), np.array([-1], dtype=int))
        for x in result:
            x.setflags(write=False)
        self._residuals[key] = result
        return result

    def _metadata(self, origin: int, indices: np.ndarray, method: str) -> dict:
        return {'origin': origin, 'origin_date': (START_DATE + timedelta(days=origin)).isoformat(), 'latest_actual_day_index': origin - 1, 'historical_indices': indices.tolist(), 'method': method, 'cold_start': self.cold_start}

    def scenarios(self, origin: int, K: int = 30, method: str = "residual") -> tuple[np.ndarray, np.ndarray, dict]:
        origin = self._origin(origin)
        lh, ph = self.point(origin)
        e_load, e_pv, indices = self.residuals(origin, K)
        metadata = self._metadata(origin, indices, method)
        if method == "raw" and origin:
            loads, pvs = self.data.load[indices].copy(), self.data.pv[indices].copy()
        elif method in ("residual", "shuffled", "raw"):
            loads, pvs = np.maximum(lh[None, :] + e_load, 0), np.maximum(ph[None, :] + e_pv, 0)
            pvs[:, ph <= 1e-9] = 0.0
            if method == "shuffled":
                # Permute scenario labels independently at each slot; retain
                # each slot's joint L/PV marginal while destroying path links.
                rng = np.random.default_rng(1729 + origin)
                for t in range(T):
                    order = rng.permutation(len(loads))
                    loads[:, t] = loads[order, t]
                    pvs[:, t] = pvs[order, t]
                metadata["shuffle_seed"] = 1729 + origin
                metadata["shuffle_scope"] = "independent scenario permutation by slot; paired L/PV"
        else:
            raise ValueError("method must be residual, raw, or shuffled")
        return loads, pvs, metadata

    def _conditional_base(self, origin: int, K: int, method: str) -> tuple:
        key = (origin, int(K), method)
        if key in self._conditional_bases:
            return self._conditional_bases[key]
        lh, ph = self.point(origin)
        e_load, e_pv, indices = self.residuals(origin, 30)
        if method != "residual":
            loads, pvs, _ = self.scenarios(origin, 30, method)
            e_load, e_pv = loads - lh[None, :], pvs - ph[None, :]
        # A shuffled-path negative control must not inherit temporal
        # persistence from the original residual archive. Fit rho to the
        # actual path representation being supplied to this controller.
        rho_load, rho_pv = ar1_coef(e_load), ar1_coef(e_pv)
        order = np.arange(len(e_load))
        if len(order) > K:
            order = np.random.default_rng(origin).choice(len(order), K, replace=False)
        result = (lh, ph, e_load[order], e_pv[order], rho_load, rho_pv, indices[order])
        self._conditional_bases[key] = result
        return result

    def conditional(self, origin: int, t: int, observed_load: np.ndarray,
                    observed_pv: np.ndarray, observe_current: bool = True,
                    K: int = 10, lag_anchor: str = "reference",
                    method: str = "residual") -> tuple[np.ndarray, np.ndarray, dict]:
        """Remaining-day conditional scenarios from an explicitly bounded prefix."""
        origin = self._origin(origin)
        if not isinstance(t, (int, np.integer)) or not 0 <= t < T or K < 1:
            raise ValueError("t must be an interval index and K must be positive")
        if lag_anchor not in ("reference", "corrected"):
            raise ValueError("lag_anchor must be reference or corrected")
        allowed = int(t) + int(bool(observe_current))
        observed_load = np.asarray(observed_load, dtype=float)
        observed_pv = np.asarray(observed_pv, dtype=float)
        if observed_load.shape != (allowed,) or observed_pv.shape != (allowed,):
            raise ValueError(f"Only an observation prefix of exactly {allowed} slots is allowed")
        if not np.isfinite(observed_load).all() or not np.isfinite(observed_pv).all():
            raise ValueError("Observation prefix contains nonfinite values")
        lh, ph, e_load, e_pv, rho_load, rho_pv, indices = self._conditional_base(origin, K, method)
        metadata = self._metadata(origin, indices, method)
        metadata.update(t=int(t), observe_current=bool(observe_current), last_observed_slot=allowed - 1, lag_anchor=lag_anchor, rho_load=rho_load, rho_pv=rho_pv, scenario_seed=origin)
        if method == "shuffled":
            metadata["shuffle_seed"] = 1729 + origin
        if not observe_current and t == 0 and lag_anchor == "corrected":
            loads = np.maximum(lh[None, :] + e_load, 0)
            pvs = np.maximum(ph[None, :] + e_pv, 0)
            metadata["residual_anchor_slot"] = None
        else:
            anchor_observed = t if observe_current else t - 1
            err_load = observed_load[anchor_observed] - lh[anchor_observed] if allowed else 0.0
            err_pv = observed_pv[anchor_observed] - ph[anchor_observed] if allowed else 0.0
            steps = np.arange(T - t) + (0 if observe_current else 1)
            factor_load, factor_pv = rho_load ** steps, rho_pv ** steps
            anchor_residual = t if observe_current or lag_anchor == "reference" else t - 1
            loads = np.maximum(lh[t:][None, :] + factor_load[None, :] * err_load
                               + e_load[:, t:] - factor_load[None, :] * e_load[:, [anchor_residual]], 0)
            pvs = np.maximum(ph[t:][None, :] + factor_pv[None, :] * err_pv
                             + e_pv[:, t:] - factor_pv[None, :] * e_pv[:, [anchor_residual]], 0)
            metadata["residual_anchor_slot"] = int(anchor_residual)
        pvs[:, ph[t:] <= 1e-9] = 0.0
        if observe_current:
            loads[:, 0] = observed_load[-1]
            pvs[:, 0] = observed_pv[-1]
        return loads, pvs, metadata
