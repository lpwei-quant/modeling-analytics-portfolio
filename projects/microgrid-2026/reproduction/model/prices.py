"""Immutable attachment-4 prices and origin-specific available price profiles."""
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

MODES = ('known_day_ahead', 'previous_day_forecast')

@dataclass(frozen=True)
class PriceDay:
    day: int
    mode: str
    source_day: int
    decision: np.ndarray

    def __post_init__(self):
        p = np.asarray(self.decision, float).copy()
        if p.shape != (144,) or not np.isfinite(p).all() or p.min() <= 0:
            raise ValueError('Decision prices must be positive finite 144-slot values')
        if self.mode not in MODES or self.source_day != self.day - int(self.mode == 'previous_day_forecast'):
            raise ValueError('Price origin violates the declared information mode')
        p.setflags(write=False)
        object.__setattr__(self, 'decision', p)

    @property
    def next_day_forecast(self):
        # A declared persistence forecast, never the next observed price row.
        return self.decision

    def metadata(self):
        return dict(mode=self.mode, origin_day=self.day, decision_price_source_day=self.source_day,
            nextday_price_source_day=self.source_day,
            nextday_method='persistence of the current origin decision-price profile',
            current_full_price_profile_known_day_ahead_assumed=self.mode == 'known_day_ahead',
            price_profile_frozen_for_intraday_control=True,
            actual_nextday_price_not_read=True, price_unit='yuan/kWh')

@dataclass(frozen=True)
class PriceArchive:
    actual: np.ndarray
    dates: tuple
    source: str

    def __post_init__(self):
        p = np.asarray(self.actual, float).copy()
        if p.shape != (365, 144) or not np.isfinite(p).all() or p.min() <= 0 or len(self.dates) != 365:
            raise ValueError('Attachment-4 prices must be positive finite 365 x 144')
        p.setflags(write=False)
        object.__setattr__(self, 'actual', p)

    def at_origin(self, day, mode):
        if mode not in MODES or not 0 <= day < 365:
            raise ValueError('Invalid origin/mode')
        source_day = day - int(mode == 'previous_day_forecast')
        if source_day < 0:
            raise ValueError('Previous-day price is unavailable on source-year day zero')
        return PriceDay(day, mode, source_day, self.actual[source_day])

    def settle_price(self, day):
        return self.actual[day]

def load_prices(root, expected_dates):
    path = Path(root) / 'problem/attachments/C题/附件4.xlsx'
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if book.sheetnames != ['Sheet1']:
            raise ValueError('Unexpected attachment-4 sheets')
        rows = list(book['Sheet1'].iter_rows(values_only=True))
    finally:
        book.close()
    if len(rows) != 366 or any(len(row) != 145 for row in rows):
        raise ValueError('Unexpected attachment-4 date/slot shape')
    dates = tuple(row[0].date() if isinstance(row[0], datetime) else row[0] for row in rows[1:])
    if dates != tuple(expected_dates):
        raise ValueError('Attachment-4 dates differ from attachment-2 actual rows')
    # Check every interval-end header against attachment 2; never shift positions.
    book = load_workbook(Path(root) / 'problem/attachments/C题/附件2.xlsx', read_only=True, data_only=True)
    try:
        actual_header = next(book.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True))
    finally:
        book.close()
    if rows[0][1:] != actual_header[1:]:
        raise ValueError('Attachment-4 time labels are not position-aligned with actuals')
    return PriceArchive(np.asarray([row[1:] for row in rows[1:]], float), dates, str(path))
