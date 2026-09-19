"""Export the five official templates from freshly computed numerical ledgers."""

import argparse
from copy import copy
from datetime import date, datetime, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.workbook.properties import CalcProperties
from run import ROOT, save
CELL_XML = re.compile(rb'(<c\b[^>]*\br="([^"]+)"[^>]*>)(.*?)(</c>)', re.DOTALL)
VALUE_XML = re.compile(rb"<v(?:\s[^>]*)?>.*?</v>|<v\s*/>", re.DOTALL)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def sequential_sum(values) -> float:
    """Preserve the previous SUM cache's left-to-right binary64 addition."""
    total = 0.0
    for value in values:
        total += float(value)
    return total

def exact_numeric_xml(source: Path, output: Path, sheet_numbers: dict[str, dict[str, float]]) -> None:
    """Keep binary64 source numbers and supply known SUM caches in OOXML.

    openpyxl normally serializes 16 significant digits and cannot calculate
    formula caches. Only our explicitly written numeric cells and SUM totals
    are patched. Formula expressions remain intact and Excel can recalculate
    them after an edit. This is not a general-purpose formula evaluator.
    """
    with ZipFile(source) as original, ZipFile(output, "w") as target:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename in sheet_numbers:
                numbers = sheet_numbers[entry.filename]
                seen = set()

                def patch(match):
                    address = match.group(2).decode("ascii")
                    if address not in numbers:
                        return match.group(0)
                    if b't="inlineStr"' in match.group(1) or b't="s"' in match.group(1):
                        raise ValueError(f"Expected numeric cell: {entry.filename}!{address}")
                    seen.add(address)
                    value = float(numbers[address])
                    if not math.isfinite(value):
                        raise ValueError(f"Non-finite output: {address}")
                    text_value = str(int(value)) if value.is_integer() else repr(value)
                    replacement = b"<v>" + text_value.encode("ascii") + b"</v>"
                    body = match.group(3)
                    body = VALUE_XML.sub(replacement, body) if VALUE_XML.search(body) else body + replacement
                    return match.group(1) + body + match.group(4)

                data = CELL_XML.sub(patch, data)
                if seen != set(numbers):
                    raise ValueError(f"Numeric cells missing in saved XML: {entry.filename}")
                data = data.replace(b' t="n"', b'')  # OOXML's default numeric type.
            target.writestr(entry, data, compress_type=ZIP_DEFLATED, compresslevel=9)

def export_one(name: str, model: dict, template: Path, destination: Path) -> dict:
    workbook = load_workbook(template)
    numeric = {sheet.title: {} for sheet in workbook}

    def put(sheet, row, column, value, number_format=None):
        cell = sheet.cell(row, column, value)
        if number_format is not None:
            cell.number_format = number_format
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            numeric[sheet.title][cell.coordinate] = float(value)
        return cell

    if model["kind"] == "q1":
        if len(model["q"]) != 144 or len(model["storage"]) != 6:
            raise ValueError("Q1 result dimensions changed")
        plan, storage = workbook["计划购电量"], workbook["充放电量"]
        # Preserve every official time label, including its shifted labels.
        for i, value in enumerate(model["q"], 2):
            put(plan, i, 2, value, "0.00000")
        for i, row in enumerate(model["storage"], 2):
            for j, value in enumerate(row, 2):
                put(storage, i, j, value, "0.00000")
        put(storage, 2, 5, model["initial_soc"], "0.00000")
        put(storage, 3, 5, model["terminal_soc"], "0.00000")
        storage.column_dimensions["E"].width = 20
    else:
        dates = model["dates"]
        if len(dates) != 334 or len(model["storage_rows"]) != 2004:
            raise ValueError("A complete 334-day computed trajectory is required")
        plan_specs = [("计划购电量", "q", "plan_fee")]
        if model["kind"] == "amendment":
            plan_specs.append(("调整购电量", "m", "final_fee"))
        for sheet_name, quantity, fee in plan_specs:
            sheet = workbook[sheet_name]
            if (sheet.max_row, sheet.max_column) != (335, 147):
                raise ValueError("Official plan template dimensions changed")
            for i, (day, values, cost) in enumerate(zip(dates, model[quantity], model[fee]), 2):
                if len(values) != 144 or sheet.cell(i, 1).value.date().isoformat() != day:
                    raise ValueError(f"Date/slot mismatch: {sheet_name} row {i}")
                for j, value in enumerate(values, 2):
                    put(sheet, i, j, value, "0.00")
                total = put(sheet, i, 146, f"=SUM(B{i}:EO{i})", "0.00")
                numeric[sheet_name][total.coordinate] = sequential_sum(values)
                put(sheet, i, 147, cost, "0.00")

        storage = workbook["充放电量"]
        block_styles = [[copy(storage.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
        for i, row in enumerate(model["storage_rows"], 2):
            for j, value in enumerate(row, 1):
                cell = storage.cell(i, j)
                if not cell.has_style and model["kind"] == "no_amendment":
                    cell._style = copy(block_styles[(i - 2) % 6][j - 1])
                cell.value = None
                if j == 1 and value is not None:
                    value = datetime.fromisoformat(value)
                put(storage, i, j, value)
                if j == 1:
                    cell.number_format = "yyyy-mm-dd"
                elif j in (3, 4, 6):
                    cell.number_format = "0.00"
                elif j == 5 and (i - 2) % 6 == 0:
                    cell.number_format = "h:mm"

        emergency = workbook["紧急购电量"]
        first_styles = [copy(emergency.cell(2, c)._style) for c in range(1, 4)]
        for i, row in enumerate(model["emergency_rows"], 2):
            for j, value in enumerate(row, 1):
                cell = emergency.cell(i, j)
                if not cell.has_style and model["kind"] == "no_amendment":
                    cell._style = copy(first_styles[j - 1])
                cell.value = None
                if j == 1 and value is not None:
                    value = datetime.fromisoformat(value)
                put(emergency, i, j, value)
                if j == 1:
                    cell.number_format = "yyyy-mm-dd"
                elif j == 3:
                    cell.number_format = "0.00"

        if model["kind"] == "no_amendment":
            # These are the readability adjustments in the original Q2/4-2 builder.
            for column, width in {"A": 14, "B": 19, "C": 16, "D": 16, "F": 18}.items():
                storage.column_dimensions[column].width = width
            for column, width in {"A": 14, "B": 21, "C": 17}.items():
                emergency.column_dimensions[column].width = width
            line = Side(style="thin", color="FFD9D9D9")
            for sheet in (storage, emergency):
                for row in sheet.iter_rows(min_row=2):
                    sheet.row_dimensions[row[0].row].height = 21
                    for cell in row:
                        cell.font = Font(name="SimSun", size=11, color="FF000000")
                        cell.border = Border(left=line, right=line, top=line, bottom=line)
                        right = (sheet is storage and cell.column in (3, 4, 6)) or (sheet is emergency and cell.column == 3)
                        cell.alignment = Alignment(horizontal="right" if right else "center", vertical="center")

    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    archive_numbers = {f"xl/worksheets/sheet{i}.xml": numeric[s.title] for i, s in enumerate(workbook, 1)}
    with tempfile.TemporaryDirectory(prefix="workbook_export_", dir=destination.parent) as temporary:
        intermediate = Path(temporary) / "openpyxl.xlsx"
        finished = Path(temporary) / name
        workbook.save(intermediate)
        workbook.close()
        exact_numeric_xml(intermediate, finished, archive_numbers)
        finished.replace(destination)
    return {"name": name, "output_sha256": digest(destination), "bytes": destination.stat().st_size,
            "explicit_numeric_cells_and_formula_caches": sum(map(len, numeric.values()))}


def interval(first,last):
    return f'{first//6}:{10*(first%6):02d}-{last//6}:{10*(last%6):02d}'


def annual_input(path,amendment):
    with np.load(path,allow_pickle=False) as source:
        z = {key:source[key].copy() for key in source.files}
    assert np.array_equal(z['d'],np.arange(31,365))
    dates = [(date(2025,1,1)+timedelta(days=int(d))).isoformat() for d in z['d']]
    storage,emergency = [],[]
    for i,day in enumerate(dates):
        for block in range(6):
            section = slice(block*24,(block+1)*24)
            storage.append([day if block == 0 else None,interval(block*24,(block+1)*24),
                float(z['c'][i,section].sum()),float(z['b'][i,section].sum()),
                0 if block == 0 else '24:00' if block == 1 else None,
                float(z['soc_start'][i,0]) if block == 0 else float(z['soc_end'][i,-1]) if block == 1 else None])
        first = None; pieces = []
        for t in range(145):
            active = t < 144 and z['e'][i,t] > 1e-6
            if active and first is None: first = t
            elif not active and first is not None:
                pieces.append([day if not pieces else None,interval(first,t),math.fsum(z['e'][i,first:t])])
                first = None
        emergency.extend(pieces or [[day,'无',0.]])
    plan_fee = z['plan_cost'] if amendment else z['cash_cost']
    model = dict(kind='amendment' if amendment else 'no_amendment',dates=dates,q=z['q'].tolist(),
        plan_fee=[math.fsum(row) for row in plan_fee],storage_rows=storage,emergency_rows=emergency)
    if amendment: model.update(m=z['m'].tolist(),final_fee=[math.fsum(row) for row in z['cash_cost']])
    return model


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out',type=Path,default=ROOT/'output')
    parser.add_argument('--destination',type=Path,default=ROOT/'result_tables')
    args = parser.parse_args();args.destination.mkdir(parents=True,exist_ok=True)
    with np.load(args.out/'q1/base/solution.npz') as z:
        model = dict(kind='q1',q=z['g'].tolist(),initial_soc=float(z['soc'][0]),
            terminal_soc=float(z['soc'][-1]),storage=[[float(z[k][b*24:(b+1)*24].sum())
            for k in ('c','d')] for b in range(6)])
    models = {'result1.xlsx':model}
    selected = json.loads((args.out/'q3/selection.json').read_text('utf-8'))['selected']
    for name,case,amend in [('result2.xlsx','q2/F_fast',False),('result3.xlsx',selected,True),
        ('result4-2.xlsx','q4_2/known_day_ahead',False),('result4-3.xlsx','q4_3/known_day_ahead',True)]:
        models[name] = annual_input(args.out/case/'ledger.npz',amend)
    report = [export_one(name,model,ROOT/'templates'/name,args.destination/name) for name,model in models.items()]
    save(args.destination/'export_report.json',dict(source='Freshly computed output ledgers',workbooks=report))
    print(json.dumps(dict(exported=len(report),destination=str(args.destination)),ensure_ascii=False))


if __name__ == '__main__':
    main()
